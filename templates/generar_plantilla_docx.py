#!/usr/bin/env python3
"""
Genera templates/plantilla_A4_horizontal.docx: documento Word compatible con
Google Docs, páginas A4 horizontales, con la cabecera v2 (sin la línea
"Marca rellenando…") en la sección de encabezado, escalada al ancho útil de la
hoja manteniendo sus proporciones. Márgenes mínimos (4 mm arriba y a los
lados) para que los marcadores ArUco no caigan en la zona que las impresoras
no imprimen; el diseño tiene además 3,6 mm de blanco propios, así que los
marcadores quedan a unos 9 mm del borde del papel.

Uso: python3 templates/generar_plantilla_docx.py <cabecera_600dpi.png> <salida.docx> [plantilla_base.pdf]
  - cabecera_600dpi.png: cabecera v2 recortada (210 x 25 mm) a 600 dpi, p. ej.
      pdftoppm -r 600 -png -singlefile -x 0 -y 0 -W 4961 -H 591 templates/cabecera_v2.pdf cabecera_600
  - plantilla_base.pdf (opcional): PDF del que copiar las preguntas al cuerpo.
"""
import sys, re
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

PNG, OUT = sys.argv[1], sys.argv[2]
BASE_PDF = sys.argv[3] if len(sys.argv) > 3 else None

MARGEN_SUP_MM, MARGEN_LAT_MM, MARGEN_INF_MM = 4, 4, 10
ANCHO_MM, ALTO_MM = 297, 210          # A4 horizontal

doc = Document()
# Tipografía del cuerpo: Arial (disponible en Google Docs)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(0)

sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Mm(ANCHO_MM), Mm(ALTO_MM)
sec.left_margin = sec.right_margin = Mm(MARGEN_LAT_MM)
sec.top_margin = Mm(MARGEN_SUP_MM)
sec.bottom_margin = Mm(MARGEN_INF_MM)
sec.header_distance = Mm(MARGEN_SUP_MM)
sec.footer_distance = Mm(5)

# Encabezado: la cabecera v2 como imagen a todo el ancho útil (297 - 2 x 4 mm),
# menos 1 mm por lado de holgura: LibreOffice desplaza ~3 mm las imágenes en
# línea del encabezado y así nunca invade el margen derecho.
header = sec.header
header.is_linked_to_previous = False
p = header.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(0)
p.add_run().add_picture(PNG, width=Mm(ANCHO_MM - 2 * MARGEN_LAT_MM - 2))

# Cuerpo: preguntas de la plantilla base (si se indica) como párrafos editables
parrafos = []
if BASE_PDF:
    import pdfplumber
    PT = 25.4 / 72
    with pdfplumber.open(BASE_PDF) as pdf:
        pg = pdf.pages[0]
        lineas = {}
        for w in pg.extract_words(use_text_flow=True):
            if w['top'] * PT >= 30:
                lineas.setdefault(round(w['top'] * PT, 1), []).append(w)
        for y in sorted(lineas):
            texto = ' '.join(w['text'] for w in sorted(lineas[y], key=lambda w: w['x0']))
            if re.match(r'^(\d+:|[a-d]\))', texto) or not parrafos:
                parrafos.append(texto)
            else:
                parrafos[-1] += ' ' + texto
for texto in parrafos:
    par = doc.add_paragraph(texto)
    if re.match(r'^\d+:', texto):
        par.paragraph_format.space_before = Pt(10)
        par.paragraph_format.keep_with_next = True
if not parrafos:
    doc.add_paragraph('')

doc.core_properties.title = 'Plantilla OMR A4 horizontal (cabecera v2)'
doc.core_properties.author = ''
doc.save(OUT)
print('escrito:', OUT, '| párrafos en el cuerpo:', len(parrafos))
